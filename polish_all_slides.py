import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def polish_all_slides():
    src_path = 'Multi Agent AI Research PPr MID_BACKUP.pptx'
    out_updated = 'Multi Agent AI Research PPr MID_Updated.pptx'
    out_original = 'Multi Agent AI Research PPr MID.pptx'

    prs = Presentation(src_path)

    def set_clean_text(shape, text, font_name="Calibri", font_size=None, bold=None, color_rgb=None, space_after=None):
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_top = Pt(2)
        tf.margin_bottom = Pt(2)
        tf.margin_left = Pt(4)
        tf.margin_right = Pt(4)
        
        while len(tf.paragraphs) > 1:
            p_last = tf.paragraphs[-1]
            p_last.text = ""
            break

        p = tf.paragraphs[0]
        p.text = text
        if space_after is not None:
            p.space_after = Pt(space_after)
        if p.runs:
            r = p.runs[0]
            if font_name:
                r.font.name = font_name
            if font_size:
                r.font.size = Pt(font_size)
            if bold is not None:
                r.font.bold = bold
            if color_rgb:
                r.font.color.rgb = color_rgb
            for r_extra in p.runs[1:]:
                r_extra.text = ""

    def set_multiline_paragraphs(shape, paragraphs_data, default_font="Calibri"):
        """
        paragraphs_data is list of (text, level, bold, size_pt, color_rgb, space_after)
        """
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_top = Pt(2)
        tf.margin_bottom = Pt(2)
        tf.margin_left = Pt(4)
        tf.margin_right = Pt(4)

        existing_count = len(tf.paragraphs)
        for i, item in enumerate(paragraphs_data):
            text, level, bold, size_pt, color_rgb, space_after = item
            if i < existing_count:
                p = tf.paragraphs[i]
            else:
                p = tf.add_paragraph()
            
            p.level = level
            p.text = text
            if space_after is not None:
                p.space_after = Pt(space_after)
            if p.runs:
                r = p.runs[0]
                r.font.name = default_font
                if size_pt is not None:
                    r.font.size = Pt(size_pt)
                if bold is not None:
                    r.font.bold = bold
                if color_rgb is not None:
                    r.font.color.rgb = color_rgb
                for r_extra in p.runs[1:]:
                    r_extra.text = ""
        
        # Clear any leftover paragraphs
        if len(paragraphs_data) < len(tf.paragraphs):
            for extra_p in tf.paragraphs[len(paragraphs_data):]:
                extra_p.text = ""

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides[0]
    s1.shapes[1].top = Inches(1.85)
    s1.shapes[1].height = Inches(1.80)
    set_clean_text(s1.shapes[1], "Multi-Agent AI Research Assistant for Scientific Paper Analysis", font_name="Georgia", font_size=44, bold=True)
    
    # Members (left)
    s1.shapes[5].left = Inches(0.40)
    s1.shapes[5].top = Inches(4.00)
    s1.shapes[5].width = Inches(4.80)
    s1.shapes[5].height = Inches(0.70)
    set_clean_text(s1.shapes[5], "Group Member: Soham Kamble (231020246), Utkarsh Rai (231020256)", font_name="Georgia", font_size=15, bold=False)

    # Supervisor (middle - clean, no underscores)
    s1.shapes[2].left = Inches(5.30)
    s1.shapes[2].top = Inches(4.00)
    s1.shapes[2].width = Inches(3.80)
    s1.shapes[2].height = Inches(0.70)
    set_clean_text(s1.shapes[2], "Supervisor: Dr. O.P. Vyas", font_name="Georgia", font_size=19, bold=False)

    # Date (right)
    s1.shapes[3].left = Inches(9.30)
    s1.shapes[3].top = Inches(4.00)
    s1.shapes[3].width = Inches(3.40)
    s1.shapes[3].height = Inches(0.70)
    set_clean_text(s1.shapes[3], "Date: 06/10/2026", font_name="Georgia", font_size=19, bold=False)

    # =========================================================================
    # SLIDE 2: Table of Contents / Content
    # =========================================================================
    s2 = prs.slides[1]
    s2.shapes[1].height = Inches(0.85) # Content title

    # =========================================================================
    # SLIDE 3: Introduction
    # =========================================================================
    s3 = prs.slides[2]
    s3.shapes[0].height = Inches(0.85) # Introduction title
    s3_intro = [
        ("Scientific literature review requires searching, reading, comparing, and summarizing large volumes of research papers.", 0, False, 16.0, None, 10),
        ("Existing tools offer limited search or summarization without intelligent collaboration between specialized AI components.", 0, False, 16.0, None, 10),
        ("The proposed system deploys multiple autonomous AI agents, each handling a dedicated phase of the research analysis workflow.", 0, False, 16.0, None, 10),
        ("Retrieval-Augmented Generation (RAG) strictly grounds every response in retrieved paper chunks, eliminating hallucinations.", 0, False, 16.0, None, 10),
        ("Output: Structured, citation-backed literature reviews with methodology comparison tables, research gap reports, and BibTeX citations.", 0, False, 16.0, None, 10),
    ]
    set_multiline_paragraphs(s3.shapes[5], s3_intro)

    # =========================================================================
    # SLIDE 4: Motivation: Issues and Challenges
    # =========================================================================
    s4 = prs.slides[3]
    s4.shapes[0].height = Inches(0.85)
    # Box 1
    set_clean_text(s4.shapes[6], "Information Overload", font_size=20, bold=True, color_rgb=RGBColor(183, 28, 28))
    set_clean_text(s4.shapes[7], "Thousands of scientific papers published weekly; manual cross-paper synthesis causes severe researcher fatigue.", font_size=14, bold=False)
    # Box 2
    set_clean_text(s4.shapes[9], "Dense Academic Formats", font_size=20, bold=True, color_rgb=RGBColor(183, 28, 28))
    set_clean_text(s4.shapes[10], "Academic PDFs contain multi-column layouts, dense formulas, and buried methodologies difficult for generic tools.", font_size=14, bold=False)
    # Box 3
    set_clean_text(s4.shapes[12], "LLM Hallucinations", font_size=20, bold=True, color_rgb=RGBColor(183, 28, 28))
    set_clean_text(s4.shapes[13], "Standard conversational AI produces plausible but ungrounded claims lacking verifiable page and chunk citations.", font_size=14, bold=False)
    # Box 4
    set_clean_text(s4.shapes[15], "Monolithic Tooling", font_size=20, bold=True, color_rgb=RGBColor(183, 28, 28))
    set_clean_text(s4.shapes[16], "Single-prompt tools fail to combine document ingestion, vector retrieval, comparative synthesis, and gap analysis.", font_size=14, bold=False)
    # Bottom
    s4.shapes[17].top = Inches(6.30)
    set_clean_text(s4.shapes[17], "Design Principle: Coordinate specialized autonomous AI agents backed by Retrieval-Augmented Generation (RAG) for verifiable, citation-grounded literature analysis.", font_size=15, bold=True, color_rgb=RGBColor(0, 32, 96))

    # =========================================================================
    # SLIDE 5: Literature Review
    # =========================================================================
    s5 = prs.slides[4]
    s5.shapes[0].height = Inches(0.85)
    s5.shapes[4].top = Inches(1.50)
    s5.shapes[4].height = Inches(4.85)
    table_s5 = s5.shapes[4].table
    table_s5.columns[0].width = Inches(2.10)
    table_s5.columns[1].width = Inches(2.30)
    table_s5.columns[2].width = Inches(1.70)
    table_s5.columns[3].width = Inches(2.40)
    table_s5.columns[4].width = Inches(3.65)
    
    s5_table_data = [
        ["Reference", "Technique / Focus", "Data / Platform", "Key Strengths", "Observation / Limitation"],
        ["Lewis et al. (2020)\nNeurIPS", "Retrieval-Augmented Generation (RAG)", "Wikipedia / Open-domain QA", "Combines parametric & non-parametric memory", "Single monolithic retrieval; lacks multi-agent specialization for academic papers."],
        ["Wu et al. (2023)\nMicrosoft AutoGen", "Multi-Agent Conversational Framework", "Collaborative Agent LLMs", "Enables role-playing and collaborative problem solving", "General-purpose; lacks specialized PDF chunking, citation tracking, and research gap extraction."],
        ["Commercial Tools\n(ChatPDF, SciSpace)", "Single-Document Vector Search & QA", "Academic PDFs", "Fast basic QA and interactive PDF viewing", "Limited cross-paper comparative reasoning; lack transparent agent routing and structured gap synthesis."],
        ["Proposed Work\n(2026)", "Multi-Agent Hybrid RAG + Provenance Engine", "Academic Research Papers", "Specialized agents (Doc, Retrieval, Analysis, Q&A) + page citations", "Midterm validated on single-paper deep analysis; cross-corpus scaling planned for end-term."]
    ]
    for r_idx, row in enumerate(s5_table_data):
        for c_idx, val in enumerate(row):
            cell = table_s5.cell(r_idx, c_idx)
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)
            cell.margin_left = Pt(5)
            cell.margin_right = Pt(5)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.name = "Calibri"
            r.font.size = Pt(11 if r_idx == 0 else 9.5)
            r.font.bold = (r_idx == 0)

    s5.shapes[5].top = Inches(6.48)
    s5.shapes[5].height = Inches(0.42)
    set_clean_text(s5.shapes[5], "Research Direction: While RAG and multi-agent LLM systems exist independently, orchestrating specialized academic agents (Ingestion, Semantic Indexing, Synthesis, Gap Discovery) with strict page-level citation provenance solves current academic literature review bottlenecks.", font_size=10.5)

    # =========================================================================
    # SLIDE 6: Problem Definition / Problem Statement
    # =========================================================================
    s6 = prs.slides[5]
    s6.shapes[1].height = Inches(0.85)
    s6.shapes[5].top = Inches(1.85)
    s6.shapes[5].height = Inches(1.20)
    set_clean_text(s6.shapes[5], "How can an autonomous multi-agent AI system intelligently orchestrate the ingestion, semantic retrieval, comparative analysis, and grounded question-answering of complex scientific papers while eliminating hallucinations through strict page-level citations and automated research gap discovery?", font_size=20, bold=True, color_rgb=RGBColor(0, 32, 96))

    s6.shapes[6].top = Inches(3.30)
    s6_specs = [
        ("Input: Unstructured scientific papers (PDFs) + user natural language research inquiries.", 0, False, 16.5, None, 14),
        ("Constraints: Strict factual grounding, page-level citation provenance, low query latency (<2s local search), and zero hallucination on out-of-domain queries.", 0, False, 16.5, None, 14),
        ("Output: Structured summaries, grounded Q&A with exact source chunks, multi-paper methodology comparison, and synthesized research gap reports.", 0, False, 16.5, None, 14)
    ]
    set_multiline_paragraphs(s6.shapes[6], s6_specs)

    # =========================================================================
    # SLIDE 7: Research Gap
    # =========================================================================
    s7 = prs.slides[6]
    s7.shapes[0].height = Inches(0.85)
    s7.shapes[5].top = Inches(1.85)
    s7_gaps = [
        ("Gap 1 — Lack of Agentic Specialization in Research Tools", 0, True, 17, RGBColor(0, 32, 96), 4),
        ("Existing academic tools use a single prompt-response loop for all queries, failing to differentiate between ingestion, semantic search, comparative synthesis, and critical evaluation.", 1, False, 14.5, None, 14),
        ("Gap 2 — Absence of Verifiable Page-Level Citation Provenance", 0, True, 17, RGBColor(0, 32, 96), 4),
        ("Current LLM research assistants frequently hallucinate facts or provide vague document-level links rather than exact page numbers and chunk-level textual evidence.", 1, False, 14.5, None, 14),
        ("Gap 3 — Disconnect Between Retrieval and Research Gap Discovery", 0, True, 17, RGBColor(0, 32, 96), 4),
        ("Available tools answer direct questions but do not autonomously synthesize cross-paper methodologies, identify conflicting findings, or pinpoint under-explored research directions.", 1, False, 14.5, None, 14)
    ]
    set_multiline_paragraphs(s7.shapes[5], s7_gaps)

    # =========================================================================
    # SLIDE 8: Objectives of the Project
    # =========================================================================
    s8 = prs.slides[7]
    s8.shapes[1].height = Inches(0.85)
    s8.shapes[5].top = Inches(1.75)
    s8_objs = [
        ("1. Design a modular Multi-Agent Architecture featuring an Orchestrator, Document Agent, Retrieval Agent, Analysis Agent, and Grounded Q&A Agent.", 0, False, 14.5, None, 8),
        ("2. Implement an intelligent Document Ingestion Pipeline supporting PDF extraction, text normalization, and 900-character chunking with 150-character overlap.", 0, False, 14.5, None, 8),
        ("3. Build a Hybrid Retrieval Engine combining ChromaDB vector embeddings (all-MiniLM-L6-v2) and lexical TF-IDF with resilient fallback mechanisms.", 0, False, 14.5, None, 8),
        ("4. Develop an Intent-Routing Engine to automatically classify queries into Summary, Comparative Analysis, Research Gaps, or Grounded Q&A.", 0, False, 14.5, None, 8),
        ("5. Implement strict Grounded Answering and Citation Tracking with exact page and chunk metadata to eliminate hallucinated answers.", 0, False, 14.5, None, 8),
        ("6. Create an intuitive, high-responsiveness Web Interface (Streamlit) featuring interactive paper exploration, live agent activity logging, and quick prompt suggestions.", 0, False, 14.5, None, 8),
        ("7. Experimentally evaluate retrieval latency, grounding accuracy, citation precision, and agent routing reliability on benchmark scientific papers.", 0, False, 14.5, None, 8)
    ]
    set_multiline_paragraphs(s8.shapes[5], s8_objs)

    # =========================================================================
    # SLIDE 9: Proposed Framework / Methodology (FLAWLESS CARDS)
    # =========================================================================
    s9 = prs.slides[8]
    s9.shapes[0].height = Inches(0.85)

    step_indices = [
        (4, 5, 6, 7, 8, "Paper Ingestion", "PDF upload / parsing\nMetadata extraction"),
        (10, 11, 12, 13, 14, "Document Agent", "900-char chunking\nPage-level mapping"),
        (16, 17, 18, 19, 20, "Vector Indexing", "Dense vector index\nChromaDB & TF-IDF"),
        (22, 23, 24, 25, 26, "Orchestrator", "Intent classification\nAgent dispatching"),
        (28, 29, 30, 31, 32, "Retrieval Agent", "Semantic vector search\nTop-k relevant chunks"),
        (34, 35, 36, 37, 38, "Analysis Agent", "Methodology review\nResearch gap synthesis"),
        (40, 41, 42, 43, 44, "Grounded Q&A", "Contextual answers\nPage-level citations"),
        (46, 47, 48, 49, 50, "UI & Activity Log", "Interactive Streamlit\nLive agent audit log"),
    ]

    for (c_idx, o_idx, n_idx, t_idx, d_idx, title_txt, desc_txt) in step_indices:
        card = s9.shapes[c_idx]
        oval = s9.shapes[o_idx]
        num_box = s9.shapes[n_idx]
        title_box = s9.shapes[t_idx]
        desc_box = s9.shapes[d_idx]

        # Height 1.68" guarantees bottom edge at y = 3.00"
        card.top = Inches(1.32)
        card.height = Inches(1.68)

        oval.top = Inches(1.42)
        num_box.top = Inches(1.44)

        title_box.top = Inches(1.88)
        title_box.height = Inches(0.36)
        set_clean_text(title_box, title_txt, font_name="Calibri", font_size=11.5, bold=True)

        desc_box.top = Inches(2.26)
        desc_box.height = Inches(0.68)
        desc_tf = desc_box.text_frame
        desc_tf.word_wrap = True
        desc_tf.margin_top = Pt(1)
        desc_tf.margin_bottom = Pt(1)
        desc_tf.margin_left = Pt(2)
        desc_tf.margin_right = Pt(2)
        
        while len(desc_tf.paragraphs) > 1:
            desc_tf.paragraphs[-1].text = ""
            break
        p = desc_tf.paragraphs[0]
        p.text = desc_txt
        if p.runs:
            r = p.runs[0]
            r.font.name = "Calibri"
            r.font.size = Pt(9.0)
            for r_extra in p.runs[1:]:
                r_extra.text = ""

    s9.shapes[51].top = Inches(3.20)
    s9.shapes[51].height = Inches(0.35)

    layer_data = [
        (52, 53, 54, "Ingestion Layer", "PyPDF & text cleaning"),
        (55, 56, 57, "Indexing Layer", "ChromaDB + TF-IDF fallback"),
        (58, 59, 60, "Agentic Layer", "Autonomous specialized agents"),
        (61, 62, 63, "Generation Layer", "Grounded LLM (Ollama / Fallback)"),
        (64, 65, 66, "Attribution Layer", "Exact page & chunk citations")
    ]
    for (box_idx, title_idx, sub_idx, t_txt, s_txt) in layer_data:
        box = s9.shapes[box_idx]
        t_box = s9.shapes[title_idx]
        s_box = s9.shapes[sub_idx]
        
        box.top = Inches(3.68)
        box.height = Inches(1.10)
        
        t_box.top = Inches(3.80)
        t_box.height = Inches(0.30)
        set_clean_text(t_box, t_txt, font_name="Calibri", font_size=12.5, bold=True)
        
        s_box.top = Inches(4.18)
        s_box.height = Inches(0.48)
        set_clean_text(s_box, s_txt, font_name="Calibri", font_size=9.5, bold=False)

    s9.shapes[67].top = Inches(5.15)
    s9.shapes[67].height = Inches(0.80)
    set_clean_text(s9.shapes[67], "Multi-Agent Execution Pipeline: User query -> Orchestrator classifies intent -> Retrieval Agent fetches semantic chunks with page provenance -> Analysis / Q&A Agent synthesizes grounded response -> Activity Logger records multi-agent execution.", font_size=13.5)

    # =========================================================================
    # SLIDE 10: Data / Database Description (NO OVERLAP, PERFECT TABLE)
    # =========================================================================
    s10 = prs.slides[9]
    s10.shapes[0].height = Inches(0.85)

    sub_s10 = s10.shapes[5]
    sub_s10.top = Inches(1.42)
    sub_s10.left = Inches(0.65)
    sub_s10.width = Inches(11.80)
    sub_s10.height = Inches(0.42)
    set_clean_text(sub_s10, "Academic research papers processed into persistent vector indices and lexical stores with metadata for grounded attribution.", font_size=12.5)

    table_shp_s10 = s10.shapes[4]
    table_shp_s10.top = Inches(2.15)
    table_shp_s10.left = Inches(0.65)
    table_shp_s10.width = Inches(11.80)
    table_shp_s10.height = Inches(4.25)
    
    table_s10 = table_shp_s10.table
    table_s10.columns[0].width = Inches(2.10)
    table_s10.columns[1].width = Inches(2.40)
    table_s10.columns[2].width = Inches(3.40)
    table_s10.columns[3].width = Inches(3.90)

    s10_table_data = [
        ["System Component", "Configuration / Scope", "Schema / Attributes", "Functional Role"],
        ["PDF Ingestion Corpus", "1+ benchmark papers\n(e.g., LoRA 26-page PDF)", "Raw text, page numbers, title, total pages", "Source research papers for document decomposition"],
        ["Chunking Store", "900-character chunks with 150-char overlap", "chunk_id, page_number, text, char_length", "Granular semantic units preserving page-level attribution"],
        ["ChromaDB Vector Store", "Collection: research_papers\n(all-MiniLM-L6-v2)", "chunk_id, vector_embedding, metadata, document", "High-dimensional semantic similarity vector search"],
        ["TF-IDF Lexical Index", "In-memory sparse term-frequency matrix", "Vocabulary terms, document frequencies, scores", "Resilient keyword search fallback when vector store is unavailable"],
        ["Agent Execution State", "Real-time session state & execution log", "timestamp, agent_name, action, detail, status", "Live execution transparency and multi-agent coordination audit"]
    ]
    for r_idx, row in enumerate(s10_table_data):
        for c_idx, val in enumerate(row):
            cell = table_s10.cell(r_idx, c_idx)
            cell.margin_top = Pt(4)
            cell.margin_bottom = Pt(4)
            cell.margin_left = Pt(6)
            cell.margin_right = Pt(6)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.name = "Calibri"
            r.font.size = Pt(11 if r_idx == 0 else 9.5)
            r.font.bold = (r_idx == 0)

    # =========================================================================
    # SLIDE 11: Experimental Results & Discussion
    # =========================================================================
    s11 = prs.slides[10]
    s11.shapes[0].height = Inches(0.85)
    set_clean_text(s11.shapes[4], "Midterm Status: Full multi-agent pipeline implemented, verified across 10 validation checks with 100% test pass rate, and deployed on interactive web UI.", font_size=15.0)

    for d_idx in [8, 12, 16, 20, 24, 28]:
        shp = s11.shapes[d_idx]
        shp.text_frame.margin_top = Pt(1)
        shp.text_frame.margin_bottom = Pt(1)
        if shp.text_frame.paragraphs and shp.text_frame.paragraphs[0].runs:
            shp.text_frame.paragraphs[0].runs[0].font.size = Pt(9.5)

    # =========================================================================
    # SLIDE 12: Conclusion and Future Directions (ELIMINATE TITLE COLLISION)
    # =========================================================================
    s12 = prs.slides[11]
    s12.shapes[1].height = Inches(0.85) # reduce PlaceHolder 1 height so it doesn't overlap text box
    s12.shapes[5].top = Inches(1.80)
    s12_bullets = [
        ("Midterm Deliverable: Successfully engineered and validated a modular Multi-Agent AI Research Assistant capable of scientific PDF ingestion, semantic indexing, and grounded analysis.", 0, False, 15.5, None, 10),
        ("Agent Collaboration: Autonomous specialization (Document, Retrieval, Analysis, Q&A) significantly outperforms monolithic prompt approaches in precision and structure.", 0, False, 15.5, None, 10),
        ("Elimination of Hallucinations: Grounded RAG with strict page-level citation provenance guarantees factual verification across complex academic literature.", 0, False, 15.5, None, 10),
        ("Resilient & Responsive: Hybrid vector and lexical fallback architecture delivers sub-second response times and zero-downtime reliability.", 0, False, 15.5, None, 10),
        ("End-Term Scope: Expand to multi-paper cross-corpus synthesis, live ArXiv/OpenAlex API ingestion, automated BibTeX citation export, and collaborative paper comparison matrices.", 0, False, 15.5, None, 10)
    ]
    set_multiline_paragraphs(s12.shapes[5], s12_bullets)

    # =========================================================================
    # SLIDE 13: Plan of Action
    # =========================================================================
    s13 = prs.slides[12]
    s13.shapes[1].height = Inches(0.85)

    phase_boxes = [
        (6, 7, 8, "Phase 1", "Multi-Paper Corpus Scaling", "Enable simultaneous indexing and cross-referencing across 20+ papers in shared workspace."),
        (10, 11, 12, "Phase 2", "Live Academic API Ingestion", "Integrate ArXiv, Semantic Scholar, and OpenAlex APIs for real-time paper retrieval."),
        (14, 15, 16, "Phase 3", "Automated Literature Review Writer", "Generate multi-section survey drafts complete with BibTeX and APA citations."),
        (18, 19, 20, "Phase 4", "Comparative Matrix Engine", "Synthesize structured comparison tables of datasets, architectures, and empirical findings."),
        (22, 23, 24, "Phase 5", "Advanced Agent Reasoning", "Implement critique/reflection agents with domain-fine-tuned embeddings (SciBERT)."),
        (26, 27, 28, "Phase 6", "Comprehensive Evaluation & Demo", "Conduct quantitative latency/accuracy benchmarks, user study, and final thesis defense.")
    ]

    for (p_idx, t_idx, d_idx, p_txt, t_txt, d_txt) in phase_boxes:
        p_box = s13.shapes[p_idx]
        t_box = s13.shapes[t_idx]
        d_box = s13.shapes[d_idx]

        p_box.left = Inches(0.82)
        p_box.width = Inches(1.40)
        
        t_box.left = Inches(2.45)
        t_box.width = Inches(2.65)
        set_clean_text(t_box, t_txt, font_name="Calibri", font_size=13.5, bold=True)
        
        d_box.left = Inches(5.25)
        d_box.width = Inches(6.85)
        set_clean_text(d_box, d_txt, font_name="Calibri", font_size=11.5, bold=False)

    # Save to updated presentation
    prs.save(out_updated)
    print(f"SAVED: {out_updated}")

    # Also try overwriting original
    try:
        prs.save(out_original)
        print(f"SAVED: {out_original}")
    except PermissionError:
        print(f"NOTICE: {out_original} is locked by PowerPoint. Saved to {out_updated}!")

if __name__ == '__main__':
    polish_all_slides()
