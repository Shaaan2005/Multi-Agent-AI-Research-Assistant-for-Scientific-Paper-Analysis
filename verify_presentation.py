from pptx import Presentation

prs = Presentation('Multi Agent AI Research PPr MID.pptx')

forbidden_keywords = [
    'examination', 'question paper', 'centre', 'smart contract', 'blockchain',
    'leak tracing', 'AES-256', 'cryptographic', 'time-lock', 'custody', 'scraped', 'Kulkarni'
]

findings = []
full_dump = []

for s_idx, slide in enumerate(prs.slides, 1):
    full_dump.append(f"\n==================== SLIDE {s_idx} ====================")
    slide_text = ""
    for shp_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    full_dump.append(f"[{shp_idx}] {txt}")
                    slide_text += " " + txt
        elif shape.has_table:
            full_dump.append(f"[{shp_idx}] [TABLE]")
            for row in shape.table.rows:
                row_str = " | ".join(c.text.strip().replace('\n', ' ') for c in row.cells)
                full_dump.append(f"    {row_str}")
                slide_text += " " + row_str
                
    for kw in forbidden_keywords:
        if kw.lower() in slide_text.lower():
            findings.append((s_idx, kw))

with open('verified_slides_dump.txt', 'w', encoding='utf-8') as f:
    f.write("\n".join(full_dump))

print(f"Total slides checked: {len(prs.slides)}")
if findings:
    print("WARNING: Found leftover keywords:")
    for s_idx, kw in findings:
        print(f"  Slide {s_idx}: keyword '{kw}'")
else:
    print("PERFECT: Zero leftover keywords found! All slides cleanly customized.")
