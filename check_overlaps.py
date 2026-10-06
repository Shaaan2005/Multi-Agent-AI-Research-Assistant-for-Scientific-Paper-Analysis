from pptx import Presentation

prs = Presentation('Multi Agent AI Research PPr MID_Updated.pptx')

for s_num in range(1, 15):
    slide = prs.slides[s_num - 1]
    print(f"=== SLIDE {s_num} ===")
    shapes_info = []
    for idx, shp in enumerate(slide.shapes):
        t = shp.top / 914400.0
        l = shp.left / 914400.0
        w = shp.width / 914400.0
        h = shp.height / 914400.0
        b = t + h
        r = l + w
        txt = ""
        if shp.has_text_frame:
            txt = shp.text_frame.text.replace('\n', ' ')
        elif shp.has_table:
            txt = f"[TABLE {len(shp.table.rows)}x{len(shp.table.columns)}]"
        shapes_info.append({
            'idx': idx, 'name': shp.name, 't': t, 'l': l, 'w': w, 'h': h, 'b': b, 'r': r,
            'txt': txt.strip()
        })
    
    # Check vertical collisions between non-empty text/table shapes
    content_shapes = [s for s in shapes_info if s['txt'] and not any(k in s['name'].lower() for k in ['date', 'footer', 'slide number'])]
    for i in range(len(content_shapes)):
        for j in range(i + 1, len(content_shapes)):
            s1 = content_shapes[i]
            s2 = content_shapes[j]
            # check bounding box overlap
            h_overlap = not (s1['r'] < s2['l'] or s1['l'] > s2['r'])
            v_overlap = not (s1['b'] < s2['t'] or s1['t'] > s2['b'])
            if h_overlap and v_overlap:
                print(f"  OVERLAP DETECTED: Shape {s1['idx']} ('{s1['txt'][:30]}') AND Shape {s2['idx']} ('{s2['txt'][:30]}')")
